from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .db import ensure_checklist_items, log_activity, next_case_number, utcnow

# RMA Number is the canonical tracked identifier inside EZ Expedite.
# These aliases exist only so common source systems can be imported without
# forcing users to rename columns first.
RMA_NUMBER_ALIASES = [
    "RMA Number",
    "RMA #",
    "RMA No.",
    "RMA No",
    "RMA",
    "Quality No.",
    "Case Number",
    "Case No.",
    "Case #",
]

RMA_FIELD_ALIASES = {
    "source_type": ["Type"],
    "sales_order": ["S.O. #", "Sales Order", "Sales Order #", "SO #"],
    "job": ["Job", "Job #", "Job Number"],
    "po_receipt": ["P.O Rec #", "PO Receipt", "PO Receipt #"],
    "work_order": ["W.O.", "Work Order", "Work Order #", "WO #"],
    "rework_work_order": ["Rework W.O. #", "Rework Work Order", "Rework WO #"],
    "customer": ["Customer", "Customer Name"],
    "purchase_order": ["Customer PO#", "Customer PO #", "Purchase Order", "PO #"],
    "supplier": ["Supplier", "Vendor"],
    "part_number": ["Part Number", "Part #", "Item Number", "Item #"],
    "revision": ["Revision", "Rev"],
    "part_description": ["Description", "Part Description"],
    "quality_no": ["Quality No."],
    "customer_ncr": ["Customer NCR#", "Customer NCR #", "Customer NCR"],
    "qty_authorized": ["Qty Authorized", "Quantity Authorized"],
    "qty_received": ["Qty Received", "Quantity Received"],
    "created_date": ["Create Date", "Created Date", "Open Date"],
    "due_date": ["QC Due Date", "Due Date", "Target Date"],
    "closed_date": ["Close Date", "Closed Date"],
    "receive_date": ["Receive Date", "Received Date"],
    "department": ["Department", "Dept"],
    "owner_name": ["Responsible Empl.", "Responsible Employee", "Owner", "Assigned To"],
    "rejection_type": ["Rejection Type"],
    "discrepancy": ["Discrepancy", "Issue", "Problem Description"],
    "containment_complete": ["Containment Action Exists", "Containment Complete"],
    "corrective_action_required": ["Corrective Action Required"],
    "corrective_action_complete": ["Corrective Action Completed", "Corrective Action Complete"],
    "work_center": ["Work Center", "Work Centre"],
    "disposition_one": ["Disposition Type One", "Disposition 1"],
    "disposition_two": ["Disposition Type Two", "Disposition 2"],
    "disposition_three": ["Disposition Type Three", "Disposition 3"],
    "reject_type": ["Reject Type"],
    "total_rework_cost": ["Total Rework Cost", "Rework Cost"],
    "sales_comment": ["Sales Comment (RMA)", "RMA Comment", "RMA Comments"],
    "customer_discrepancy": ["Customer Discrepancy", "Customer Complaint Detail"],
    "status": ["Status"],
    "disposition_complete": ["Disposition Completed", "Disposition Complete"],
    "qty_returned": ["Qty Returned", "Quantity Returned"],
    "customer_complaint": ["Customer Complaint"],
}


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def yes(value: Any) -> int:
    if isinstance(value, bool):
        return int(value)
    return int(clean(value).lower() in {"1", "true", "yes", "y", "x"})


def number(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def iso_date(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (float, int)):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).date().isoformat()
        except (OverflowError, ValueError):
            return None
    text = clean(value)
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def normalize_status(value: Any) -> str:
    raw = clean(value).upper()
    if raw in {"CLOSED", "COMPLETE", "COMPLETED"}:
        return "CLOSED"
    if raw in {"OPEN", "ACTIVE"}:
        return "INVESTIGATING"
    return raw or "NEW"


def inspect_headers(path: str | Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = [clean(cell.value) for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    wb.close()
    return headers


def _first_header(index: dict[str, int], aliases: list[str]) -> str | None:
    for alias in aliases:
        if alias in index:
            return alias
    return None


def import_rma_workbook(con, path: str | Path, actor: str = "Workbook import") -> dict[str, int]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [clean(v) for v in next(rows)]

    index: dict[str, int] = {}
    for i, name in enumerate(headers):
        # Preserve the first occurrence of duplicate header names.
        index.setdefault(name, i)

    rma_header = _first_header(index, RMA_NUMBER_ALIASES)
    if not rma_header:
        wb.close()
        raise ValueError(
            "Workbook needs an RMA identifier column. Supported names include: "
            + ", ".join(RMA_NUMBER_ALIASES)
        )

    resolved: dict[str, str] = {}
    for target, aliases in RMA_FIELD_ALIASES.items():
        header = _first_header(index, aliases)
        if header:
            resolved[target] = header

    type_row = con.execute("SELECT id FROM occurrence_types WHERE name='RMA'").fetchone()
    if not type_row:
        wb.close()
        raise RuntimeError("RMA occurrence type is not configured")
    type_id = int(type_row[0])

    imported = skipped = updated = 0
    for source_row in rows:
        if not any(v not in (None, "") for v in source_row):
            skipped += 1
            continue

        rma_number = clean(source_row[index[rma_header]]) if index[rma_header] < len(source_row) else ""
        if not rma_number:
            skipped += 1
            continue

        data: dict[str, Any] = {}
        for target, header in resolved.items():
            idx = index[header]
            if idx < len(source_row):
                data[target] = source_row[idx]

        existing = con.execute(
            "SELECT occurrence_id FROM rma_details WHERE rma_number=?",
            (rma_number,),
        ).fetchone()

        title_bits = [
            clean(data.get("customer")),
            clean(data.get("part_number")),
            clean(data.get("discrepancy")),
        ]
        title = " | ".join([x for x in title_bits if x])[:240] or f"RMA {rma_number}"
        now = utcnow()
        status = normalize_status(data.get("status"))
        created_date = iso_date(data.get("created_date"))
        due_date = iso_date(data.get("due_date"))
        closed_date = iso_date(data.get("closed_date"))

        if existing:
            occurrence_id = int(existing[0])
            con.execute(
                """UPDATE occurrences SET title=?,description=?,source=?,customer=?,supplier=?,part_number=?,revision=?,
                   work_order=?,sales_order=?,purchase_order=?,department=?,work_center=?,owner_name=?,status=?,due_date=?,
                   created_date=COALESCE(?,created_date),closed_date=?,updated_at=? WHERE id=?""",
                (
                    title,
                    clean(data.get("customer_discrepancy")) or clean(data.get("discrepancy")),
                    clean(data.get("source_type")) or "RMA import",
                    clean(data.get("customer")),
                    clean(data.get("supplier")),
                    clean(data.get("part_number")),
                    clean(data.get("revision")),
                    clean(data.get("work_order")),
                    clean(data.get("sales_order")),
                    clean(data.get("purchase_order")),
                    clean(data.get("department")),
                    clean(data.get("work_center")),
                    clean(data.get("owner_name")),
                    status,
                    due_date,
                    created_date,
                    closed_date,
                    now,
                    occurrence_id,
                ),
            )
            updated += 1
        else:
            case_number = next_case_number(con, type_id)
            cur = con.execute(
                """INSERT INTO occurrences(
                   case_number,occurrence_type_id,title,description,source,customer,supplier,part_number,revision,
                   work_order,sales_order,purchase_order,department,work_center,owner_name,priority,status,due_date,
                   created_date,closed_date,last_activity_at,created_at,updated_at)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    case_number,
                    type_id,
                    title,
                    clean(data.get("customer_discrepancy")) or clean(data.get("discrepancy")),
                    clean(data.get("source_type")) or "RMA import",
                    clean(data.get("customer")),
                    clean(data.get("supplier")),
                    clean(data.get("part_number")),
                    clean(data.get("revision")),
                    clean(data.get("work_order")),
                    clean(data.get("sales_order")),
                    clean(data.get("purchase_order")),
                    clean(data.get("department")),
                    clean(data.get("work_center")),
                    clean(data.get("owner_name")),
                    "Normal",
                    status,
                    due_date,
                    created_date,
                    closed_date,
                    now,
                    now,
                    now,
                ),
            )
            occurrence_id = int(cur.lastrowid)
            imported += 1

        quality_no = clean(data.get("quality_no")) or None
        con.execute(
            """INSERT INTO rma_details(
                occurrence_id,rma_number,quality_no,customer_ncr,part_description,qty_authorized,qty_received,
                receive_date,rejection_type,discrepancy,containment_complete,corrective_action_required,
                corrective_action_complete,disposition_one,disposition_two,disposition_three,disposition_complete,
                reject_type,qty_returned,customer_complaint,total_rework_cost,recovery_status,sales_comment,
                customer_discrepancy)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                ON CONFLICT(occurrence_id) DO UPDATE SET
                rma_number=excluded.rma_number,quality_no=excluded.quality_no,customer_ncr=excluded.customer_ncr,
                part_description=excluded.part_description,qty_authorized=excluded.qty_authorized,
                qty_received=excluded.qty_received,receive_date=excluded.receive_date,
                rejection_type=excluded.rejection_type,discrepancy=excluded.discrepancy,
                containment_complete=excluded.containment_complete,
                corrective_action_required=excluded.corrective_action_required,
                corrective_action_complete=excluded.corrective_action_complete,
                disposition_one=excluded.disposition_one,disposition_two=excluded.disposition_two,
                disposition_three=excluded.disposition_three,disposition_complete=excluded.disposition_complete,
                reject_type=excluded.reject_type,qty_returned=excluded.qty_returned,
                customer_complaint=excluded.customer_complaint,total_rework_cost=excluded.total_rework_cost,
                sales_comment=excluded.sales_comment,customer_discrepancy=excluded.customer_discrepancy""",
            (
                occurrence_id,
                rma_number,
                quality_no,
                clean(data.get("customer_ncr")) or None,
                clean(data.get("part_description")) or None,
                number(data.get("qty_authorized")),
                number(data.get("qty_received")),
                iso_date(data.get("receive_date")),
                clean(data.get("rejection_type")) or None,
                clean(data.get("discrepancy")) or None,
                yes(data.get("containment_complete")),
                yes(data.get("corrective_action_required")),
                yes(data.get("corrective_action_complete")),
                clean(data.get("disposition_one")) or None,
                clean(data.get("disposition_two")) or None,
                clean(data.get("disposition_three")) or None,
                yes(data.get("disposition_complete")),
                clean(data.get("reject_type")) or None,
                number(data.get("qty_returned")),
                yes(data.get("customer_complaint")),
                number(data.get("total_rework_cost")),
                "NOT REQUIRED",
                clean(data.get("sales_comment")) or None,
                clean(data.get("customer_discrepancy")) or None,
            ),
        )
        ensure_checklist_items(con, occurrence_id, type_id)
        log_activity(con, occurrence_id, "IMPORT", f"RMA synchronized: {rma_number}.", actor)

    wb.close()
    return {"imported": imported, "updated": updated, "skipped": skipped}
